"""Branded .docx engine (python-docx), incl. teaching-companion helpers (labeled, qa).
Adapt the palette below to the brand. No em dashes in generated text."""
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.enum.text import WD_TAB_ALIGNMENT
REDc=RGBColor(0xCF,0x33,0x38); DARKc=RGBColor(0x1F,0x21,0x20); GREYc=RGBColor(0x70,0x70,0x70); RED="CF3338"
def new_doc():
    d=Document(); st=d.styles['Normal']; st.font.name='Arial'; st.font.size=Pt(11); st.font.color.rgb=DARKc
    for sec in d.sections:
        sec.top_margin=Inches(0.9); sec.bottom_margin=Inches(0.9); sec.left_margin=Inches(1.0); sec.right_margin=Inches(1.0)
    return d
def _shade(p,fill):
    pPr=p._p.get_or_add_pPr(); sh=OxmlElement('w:shd'); sh.set(qn('w:val'),'clear'); sh.set(qn('w:fill'),fill); pPr.append(sh)
def wordmark(d,right_text=None):
    p=d.add_paragraph(); p.space_after=Pt(2)
    r=p.add_run("BRAND NAME"); r.font.name='Arial'; r.bold=True; r.font.size=Pt(11); r.font.color.rgb=REDc
    if right_text:
        p.add_run("\t"); t=p.add_run(right_text); t.font.size=Pt(9); t.font.color.rgb=GREYc
        p.paragraph_format.tab_stops.add_tab_stop(Inches(6.5), WD_TAB_ALIGNMENT.RIGHT)
    pPr=p._p.get_or_add_pPr(); pbdr=OxmlElement('w:pBdr'); b=OxmlElement('w:bottom')
    b.set(qn('w:val'),'single'); b.set(qn('w:sz'),'12'); b.set(qn('w:space'),'4'); b.set(qn('w:color'),RED); pbdr.append(b); pPr.append(pbdr)
def kicker(d,text):
    p=d.add_paragraph(); p.space_after=Pt(0); r=p.add_run(text.upper()); r.bold=True; r.font.size=Pt(10.5); r.font.color.rgb=REDc; return p
def title(d,text):
    p=d.add_paragraph(); p.space_after=Pt(6); r=p.add_run(text); r.bold=True; r.font.size=Pt(22); r.font.color.rgb=DARKc; return p
def h1(d,text):
    p=d.add_paragraph(); p.space_before=Pt(14); p.space_after=Pt(4); r=p.add_run(text); r.bold=True; r.font.size=Pt(15); r.font.color.rgb=REDc; return p
def h2(d,text):
    p=d.add_paragraph(); p.space_before=Pt(12); p.space_after=Pt(2); r=p.add_run(text); r.bold=True; r.font.size=Pt(13); r.font.color.rgb=DARKc; return p
def body(d,text,italic=False,size=11):
    p=d.add_paragraph(); p.space_after=Pt(4); r=p.add_run(text); r.italic=italic; r.font.size=Pt(size); r.font.color.rgb=DARKc; return p
def bullet(d,text,bold_lead=None):
    p=d.add_paragraph(style='List Bullet'); p.space_after=Pt(2)
    if bold_lead:
        r=p.add_run(bold_lead); r.bold=True; r.font.size=Pt(11); r.font.color.rgb=DARKc
        r2=p.add_run(text); r2.font.size=Pt(11); r2.font.color.rgb=DARKc
    else:
        r=p.add_run(text); r.font.size=Pt(11); r.font.color.rgb=DARKc
    return p
def check(d,text):
    p=d.add_paragraph(); p.space_after=Pt(2); p.paragraph_format.left_indent=Inches(0.25)
    b=p.add_run("☐  "); b.font.size=Pt(11); r=p.add_run(text); r.font.size=Pt(10.5); r.font.color.rgb=DARKc; return p
def code(d,lines):
    for ln in lines:
        p=d.add_paragraph(); p.space_after=Pt(0); p.space_before=Pt(0); p.paragraph_format.left_indent=Inches(0.15)
        _shade(p,"2B2B2B"); r=p.add_run(ln if ln else " "); r.font.name='Consolas'; r.font.size=Pt(9.5); r.font.color.rgb=RGBColor(0xF5,0xF5,0xF5)
    d.add_paragraph().space_after=Pt(3)
def callout(d,label,text,fill="F3E4E4"):
    p=d.add_paragraph(); p.space_before=Pt(4); p.space_after=Pt(4); _shade(p,fill)
    p.paragraph_format.left_indent=Inches(0.1); p.paragraph_format.right_indent=Inches(0.1)
    r=p.add_run(label+"  "); r.bold=True; r.font.size=Pt(10.5); r.font.color.rgb=REDc
    r2=p.add_run(text); r2.font.size=Pt(10.5); r2.font.color.rgb=DARKc; return p
# ---- teaching-companion helpers ----
def labeled(d,label,text):   # e.g. labeled(d,"In plain English.", "...")
    p=d.add_paragraph(); p.space_after=Pt(4)
    r=p.add_run(label+"  "); r.bold=True; r.font.size=Pt(10); r.font.color.rgb=REDc
    r2=p.add_run(text); r2.font.size=Pt(11); r2.font.color.rgb=DARKc; return p
def section_label(d,text):
    p=d.add_paragraph(); r=p.add_run(text); r.bold=True; r.font.size=Pt(10); r.font.color.rgb=REDc; return p
def qa(d,q,a):   # a "They'll ask" Q/A bullet
    p=d.add_paragraph(style='List Bullet'); p.space_after=Pt(3)
    rq=p.add_run("Q: "+q+"  "); rq.bold=True; rq.font.size=Pt(10.5); rq.font.color.rgb=DARKc
    ra=p.add_run("A: "+a); ra.font.size=Pt(10.5); ra.font.color.rgb=RGBColor(0x4a,0x4a,0x4a)
def note_lines(d,n=3):
    for _ in range(n):
        p=d.add_paragraph(); p.space_after=Pt(10)
        pPr=p._p.get_or_add_pPr(); pbdr=OxmlElement('w:pBdr'); b=OxmlElement('w:bottom')
        b.set(qn('w:val'),'single'); b.set(qn('w:sz'),'4'); b.set(qn('w:space'),'2'); b.set(qn('w:color'),'CCCCCC'); pbdr.append(b); pPr.append(pbdr)

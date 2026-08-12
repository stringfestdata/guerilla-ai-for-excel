"""Spill-safe Excel / Python-in-Excel follow-along workbook engine (openpyxl).
Rules (see references/handson-file.md): nothing merged across the work column (E+),
no markers in the spill zone, wide berth after each example, code cells sized to show all code.
Adapt the palette below to the brand."""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.worksheet.properties import PageSetupProperties
from openpyxl.utils import get_column_letter
import math
RED="FFCF3338"; REDDK="FFC00000"; DARK="FF1F2120"; GREY="FF707070"
CODEBG="FFF3F1EA"; TAKEBG="FFEDE9DE"; WHITE="FFFFFFFF"; TIPBG="FFFBF3D6"; WORKBG="FFEFF5EC"
F_TITLE=Font(name="Arial Black", size=20, bold=True, color=DARK)
F_SUB=Font(name="Arial", size=11, color=GREY); F_HINTLINE=Font(name="Arial", size=10, color=GREY, italic=True)
F_STEP=Font(name="Arial", size=13, bold=True, color=WHITE); F_STEPDESC=Font(name="Arial", size=11, color=DARK)
F_LABEL=Font(name="Arial", size=10, bold=True, color=RED); F_CODE=Font(name="Consolas", size=11, color=DARK)
F_TAKE_H=Font(name="Arial", size=11, bold=True, color=RED); F_TAKE=Font(name="Arial", size=11, color=DARK)
F_WHY_H=Font(name="Arial", size=11, bold=True, color=WHITE); F_WHY=Font(name="Arial", size=10.5, color=WHITE)
F_TIP_H=Font(name="Arial", size=11, bold=True, color=DARK); F_TIP=Font(name="Arial", size=10.5, color=DARK)
F_WORK=Font(name="Arial", size=10, bold=True, color="FF2F6E2F", italic=True)
FILL_STEP=PatternFill("solid", fgColor=RED); FILL_WHY=PatternFill("solid", fgColor=REDDK)
FILL_CODE=PatternFill("solid", fgColor=CODEBG); FILL_TAKE=PatternFill("solid", fgColor=TAKEBG)
FILL_TIP=PatternFill("solid", fgColor=TIPBG); FILL_WORK=PatternFill("solid", fgColor=WORKBG)
WRAP=Alignment(wrap_text=True, vertical="top")
thin=Side(style="thin", color="FFD8D5CA"); BORDER=Border(left=thin,right=thin,top=thin,bottom=thin)
CPL=70
def _merge_bc(ws,row): ws.merge_cells(f"B{row}:C{row}")   # NEVER merge into the work column (E+)
def follow_sheet(wb, title, subtitle):
    ws=wb.active; ws.title="follow-along"; ws.sheet_properties.tabColor=RED[2:]
    ws.sheet_view.showGridLines=False
    ws.page_setup.orientation="landscape"; ws.page_setup.fitToWidth=1; ws.page_setup.fitToHeight=0
    ws.sheet_properties.pageSetUpPr=PageSetupProperties(fitToPage=True); ws.print_area="B1:K800"
    for col,w in {"A":2.5,"B":15,"C":74,"D":2}.items(): ws.column_dimensions[col].width=w
    for col in "EFGHIJK": ws.column_dimensions[col].width=13
    ws["B1"]=title; ws["B1"].font=F_TITLE
    _merge_bc(ws,2); ws["B2"]=subtitle; ws["B2"].font=F_SUB; ws["B2"].alignment=WRAP
    _merge_bc(ws,3)
    ws["B3"]=("Type each code block into a PY() cell in column E (to the right of the code) and run it. "
              "The whole area to the right and below is left open so your results can spill freely.")
    ws["B3"].font=F_HINTLINE; ws["B3"].alignment=WRAP; ws.row_dimensions[3].height=28
    wc=ws["E4"]; wc.value="↓  your PY() cells go here"; wc.font=F_WORK; wc.fill=FILL_WORK
    wc.alignment=Alignment(vertical="center", horizontal="left"); ws.row_dimensions[1].height=28
    return ws
def typealong(ws,row):
    _merge_bc(ws,row); c=ws[f"B{row}"]; c.value="Type it, don't paste it."; c.font=F_TIP_H; c.fill=FILL_TIP
    c.alignment=Alignment(vertical="center", indent=1); ws.row_dimensions[row].height=18; row+=1
    _merge_bc(ws,row); c=ws[f"B{row}"]
    c.value=("The code is printed here so you never get stuck. But time permitting, type along rather than paste. "
             "Typing builds the muscle memory and makes you slow down and understand each line. Paste only if you fall behind.")
    c.font=F_TIP; c.fill=FILL_TIP; c.alignment=WRAP; ws.row_dimensions[row].height=58; return row+2
def why_box(ws,row,header,lines):
    _merge_bc(ws,row); c=ws[f"B{row}"]; c.value=header; c.font=F_WHY_H; c.fill=FILL_WHY
    c.alignment=Alignment(wrap_text=True, vertical="center"); ws.row_dimensions[row].height=20; row+=1
    _merge_bc(ws,row); c=ws[f"B{row}"]; c.value="\n".join(lines); c.font=F_WHY; c.fill=FILL_WHY
    c.alignment=WRAP; ws.row_dimensions[row].height=16*len(lines)+8; return row+2
def tip_box(ws,row,header,text):
    _merge_bc(ws,row); c=ws[f"B{row}"]; c.value=header; c.font=F_TIP_H; c.fill=FILL_TIP
    c.alignment=Alignment(vertical="center", indent=1); ws.row_dimensions[row].height=18; row+=1
    _merge_bc(ws,row); c=ws[f"B{row}"]; c.value=text; c.font=F_TIP; c.fill=FILL_TIP; c.alignment=WRAP
    ws.row_dimensions[row].height=14*max(1,math.ceil(len(text)/95))+10; return row+2
def step(ws,row,label,desc=None):
    _merge_bc(ws,row); c=ws[f"B{row}"]; c.value=label; c.font=F_STEP; c.fill=FILL_STEP
    c.alignment=Alignment(wrap_text=True, vertical="center", indent=1); ws.row_dimensions[row].height=22; row+=1
    if desc:
        _merge_bc(ws,row); c=ws[f"B{row}"]; c.value=desc; c.font=F_STEPDESC; c.alignment=WRAP
        ws.row_dimensions[row].height=15*max(1,math.ceil(len(desc)/95))+6; row+=1
    return row
def _wl(code):
    n=0
    for ln in code.split("\n"): n+=max(1,math.ceil(len(ln)/CPL)) if ln else 1
    return n
def code_row(ws,row,label,code):
    ws[f"B{row}"]=label; ws[f"B{row}"].font=F_LABEL
    ws[f"B{row}"].alignment=Alignment(vertical="top",horizontal="right",wrap_text=True)
    cc=ws[f"C{row}"]; cc.value=code; cc.font=F_CODE; cc.fill=FILL_CODE; cc.alignment=WRAP; cc.border=BORDER
    ws.row_dimensions[row].height=max(20,15.5*_wl(code)+9); return row+1
def note_row(ws,row,text):
    _merge_bc(ws,row); c=ws[f"B{row}"]; c.value=text; c.font=F_HINTLINE; c.alignment=WRAP
    ws.row_dimensions[row].height=14*max(1,math.ceil(len(text)/95))+4; return row+1
def berth(ws,row,n):   # empty, unmerged rows so PY() output can spill without collision
    for i in range(n): ws.row_dimensions[row+i].height=15
    return row+n
def example(ws,row,label,code_lines,note=None,clear=14):
    row=code_row(ws,row,label,"\n".join(code_lines))
    if note: row=note_row(ws,row,note)
    return berth(ws,row,clear)
def takeaway(ws,row,text):
    _merge_bc(ws,row); c=ws[f"B{row}"]; c.value="✓  Takeaway"; c.font=F_TAKE_H; c.fill=FILL_TAKE
    c.alignment=Alignment(vertical="center", indent=1); ws.row_dimensions[row].height=18; row+=1
    _merge_bc(ws,row); c=ws[f"B{row}"]; c.value=text; c.font=F_TAKE; c.fill=FILL_TAKE; c.alignment=WRAP
    ws.row_dimensions[row].height=15*max(1,math.ceil(len(text)/95))+8; return row+2
def spacer(ws,row,h=10): ws.row_dimensions[row].height=h; return row+1
def data_sheet(wb,sheetname,columns,rows,tablename,tabcolor=DARK):
    ws=wb.create_sheet(sheetname); ws.sheet_properties.tabColor=tabcolor[2:]
    ws.append(list(columns))
    for r in rows: ws.append(list(r))
    ref=f"A1:{get_column_letter(len(columns))}{len(rows)+1}"
    t=Table(displayName=tablename, ref=ref); t.tableStyleInfo=TableStyleInfo(name="TableStyleMedium3", showRowStripes=True)
    ws.add_table(t)
    for i,col in enumerate(columns,1): ws.column_dimensions[get_column_letter(i)].width=max(11,min(30,len(str(col))+4))
    ws.freeze_panes="A2"; return ws

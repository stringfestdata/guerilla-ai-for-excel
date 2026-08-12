"""Build the combined Guerrilla AI cheat sheet (.docx) from content.py."""
import os, sys
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import content as C
import docbuilder as B
from docx.shared import Pt, Inches
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

OUT = sys.argv[1] if len(sys.argv) > 1 else "out"

def sf_wordmark(d, right_text):
    p = d.add_paragraph(); p.space_after = Pt(2)
    r = p.add_run("STRINGFEST ANALYTICS"); r.font.name = "Arial"; r.bold = True
    r.font.size = Pt(11); r.font.color.rgb = B.REDc
    p.add_run("\t"); t = p.add_run(right_text); t.font.size = Pt(9); t.font.color.rgb = B.GREYc
    from docx.enum.text import WD_TAB_ALIGNMENT
    p.paragraph_format.tab_stops.add_tab_stop(Inches(6.5), WD_TAB_ALIGNMENT.RIGHT)
    pPr = p._p.get_or_add_pPr(); pbdr = OxmlElement('w:pBdr'); b = OxmlElement('w:bottom')
    b.set(qn('w:val'), 'single'); b.set(qn('w:sz'), '12'); b.set(qn('w:space'), '4')
    b.set(qn('w:color'), B.RED); pbdr.append(b); pPr.append(pbdr)

def shot(d):
    B.callout(d, "Screenshot:", "[insert screenshot here]", fill="EEECE1")

summary = C.region_summary(C.clean_rows(C.junked(C.make_rows())) + C.clean_rows(C.june_rows()))

d = B.new_doc()
sf_wordmark(d, "stringfestanalytics.com")
B.kicker(d, "The five guerrilla moves, on two pages")
B.title(d, "Guerrilla AI for Excel: Cheat Sheet")
B.body(d, C.THESIS + " Here are the week's five moves, each usable the same afternoon.")

moves = [
    ("Move 1: Wake up the AI already in Excel (Day 1)",
     "Flash Fill learns a pattern from one typed example. Analyze Data answers plain-English "
     "questions with PivotTables. Both ship with Excel, no license needed.",
     [l for l in C.D1_FLASHFILL if not l.startswith("#")] +
     [l for l in C.D1_ANALYZE if not l.startswith("#")]),
    ("Move 2: Make your data AI-readable (Day 2)",
     "Every AI tool performs dramatically better on a named table than a raw range. "
     "Structured references make formulas self-documenting.",
     [l for l in C.D2_TABLE if not l.startswith("#")] +
     [l for l in C.D2_RERUN if not l.startswith("#")]),
    ("Move 3: Record the cleanup once in Power Query (Day 3)",
     "Automated, repeatable, refreshable data prep with zero AI inside. Clean inputs are "
     "what stop AI from hallucinating your numbers. Stuck on a transform? A free chatbot "
     "writes M code; paste it in the Advanced Editor.",
     [l for l in C.D3_TRANSFORMS if not l.startswith("#")] +
     [l for l in C.D3_REFRESH if not l.startswith("#")]),
    ("Move 4: Run Python behind =PY() (Day 4)",
     "A full Python runtime with pandas and seaborn sits inside Excel. Describe the analysis "
     "to a free chatbot, paste the code into a PY cell, run it on your table.",
     C.D4_BLOCK1 + C.D4_BLOCK2 + C.D4_BLOCK3),
    ("Move 5: Speak Markdown, the language AI speaks (Day 5)",
     "Structured prompts steer any model better than a blob. Tables round-trip between "
     "Excel and chatbots as Markdown, no connectors or licenses required.",
     C.d5_structured(summary) + [""] +
     [l for l in C.D5_ROUNDTRIP if not l.startswith("#")]),
]
for head, why, steps in moves:
    B.h1(d, head)
    B.body(d, why)
    B.code(d, steps)
    shot(d)

B.h1(d, "Keep going")
B.bullet(d, C.CTA_MEMBERSHIP, bold_lead="Replays: ")
B.bullet(d, C.CTA_COHORT, bold_lead="Go deeper: ")
B.bullet(d, "Questions any time: george@stringfestanalytics.com", bold_lead="Say hi: ")

p = os.path.join(OUT, "Week Assets"); os.makedirs(p, exist_ok=True)
f = os.path.join(p, "Guerrilla AI - Cheat Sheet.docx"); d.save(f); print("wrote", f)

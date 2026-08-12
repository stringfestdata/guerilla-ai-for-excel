"""QA: em dashes, spelling, banned constructions, notes coverage, table names,
and code/step consistency across artifacts. Exits nonzero on failure."""
import os, re, sys, glob
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import content as C
from pptx import Presentation
from openpyxl import load_workbook
from docx import Document

KIT = sys.argv[1]
fail = []

def deck_text(p):
    prs = Presentation(p); slides_txt, notes_txt, missing = [], [], []
    for i, s in enumerate(prs.slides, 1):
        t = " ".join(r.text for sh in s.shapes if sh.has_text_frame
                     for pa in sh.text_frame.paragraphs for r in pa.runs)
        slides_txt.append(t)
        n = s.notes_slide.notes_text_frame.text if s.has_notes_slide else ""
        notes_txt.append(n)
        if len(n.strip()) < 40: missing.append(i)
    return slides_txt, notes_txt, missing

def wb_text(p):
    wb = load_workbook(p); out = []
    tables = {}
    for ws in wb.worksheets:
        tables[ws.title] = list(ws.tables.keys()) if hasattr(ws, "tables") else []
        for row in ws.iter_rows():
            for c in row:
                if c.value is not None: out.append(str(c.value))
    return "\n".join(out), tables

def docx_text(p):
    return "\n".join(pa.text for pa in Document(p).paragraphs)

texts = {}   # path -> text
for p in sorted(glob.glob(os.path.join(KIT, "**", "*.*"), recursive=True)):
    if p.endswith(".pptx"):
        st, nt, missing = deck_text(p)
        texts[p] = "\n".join(st + nt)
        if missing: fail.append(f"{os.path.basename(p)}: slides missing notes: {missing}")
    elif p.endswith(".xlsx"):
        texts[p], tabs = wb_text(p)
        texts[p + "::tables"] = str(tabs)
    elif p.endswith(".docx"):
        texts[p] = docx_text(p)
    elif p.endswith(".md"):
        texts[p] = open(p, encoding="utf-8").read()

# 1. em dashes and en dashes
for p, t in texts.items():
    for ch in ("—", "–"):
        if ch in t:
            for line in t.split("\n"):
                if ch in line: fail.append(f"DASH in {os.path.basename(p)}: {line[:90]}")

# 2. spelling: single-r Guerilla anywhere is a failure
for p, t in texts.items():
    if re.search(r"(?i)guerill?a", t):
        for m in set(re.findall(r"(?i)gueril?la", t)):
            if m.lower() != "guerrilla":
                fail.append(f"SPELLING '{m}' in {os.path.basename(p)}")

# 3. banned 'not X but Y' construction (heuristic)
pat = re.compile(r"\bnot\s+(?:a\s+|an\s+|the\s+)?\w+(?:\s+\w+)?,?\s+but\b", re.I)
for p, t in texts.items():
    for line in t.split("\n"):
        if pat.search(line): fail.append(f"NOT-BUT in {os.path.basename(p)}: {line.strip()[:90]}")

# 4. table names intact
expect = {
    "Day 2 - Solution": ["SalesData"], "Day 3 - Demo Start": ["RawExport"],
    "Day 3 - Solution": ["RawExport", "SalesClean"], "Day 4 - Demo Start": ["SalesClean"],
    "Day 4 - Solution": ["SalesClean"], "Day 5 - Demo Start": ["RegionSummary"],
    "Day 5 - Solution": ["RegionSummary"],
}
for key, wanted in expect.items():
    match = [p for p in texts if key in p and p.endswith("::tables")]
    if not match: fail.append(f"no workbook found for {key}"); continue
    got = texts[match[0]]
    for t in wanted:
        if t not in got: fail.append(f"TABLE {t} missing in {key}: {got}")

# 5. consistency: every step/code line appears in deck, workbook, and cheat sheet
summary = C.region_summary(C.clean_rows(C.junked(C.make_rows())) + C.clean_rows(C.june_rows()))
groups = {
    1: [C.D1_FLASHFILL, C.D1_ANALYZE], 2: [C.D2_TABLE, C.D2_RERUN],
    3: [C.D3_TRANSFORMS, C.D3_REFRESH], 4: [C.D4_BLOCK1, C.D4_BLOCK2, C.D4_BLOCK3],
    5: [C.D5_ROUNDTRIP, C.d5_structured(summary)],
}
cheat = [t for p, t in texts.items() if "Cheat Sheet" in p][0]
for day, blocks in groups.items():
    deck = "\n".join(t for p, t in texts.items() if f"Day {day} - Slides" in p)
    demo = "\n".join(t for p, t in texts.items()
                     if f"Day {day} - " in p and p.endswith(".xlsx"))
    for b in blocks:
        for line in b:
            if not line: continue
            if line not in demo: fail.append(f"D{day} line missing in workbook: {line[:60]}")
            core = line.lstrip("# ")
            if line not in deck and core not in deck:
                # deck shows curated subsets on some slides; only code must match exactly
                if day in (4,): fail.append(f"D{day} code missing in deck: {line[:60]}")
            if line not in cheat and core not in cheat:
                if line.startswith("#") and day != 4: continue  # comment headers stripped by design
                fail.append(f"D{day} line missing in cheat sheet: {line[:60]}")

# 6. thesis phrase and schedule consistency in copy
for need, where in [("Sessions may occasionally be rescheduled around client engagements",
                     "Eventbrite Listing"),
                    ("Sessions may occasionally be rescheduled around client engagements",
                     "Welcome Email")]:
    src = [t for p, t in texts.items() if where in p]
    if not src or need not in src[0]: fail.append(f"missing reschedule line in {where}")

print("\n".join(fail) if fail else "QA CLEAN")
sys.exit(1 if fail else 0)

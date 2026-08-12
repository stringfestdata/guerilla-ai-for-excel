"""Build demo-start and solution workbooks for all five days from content.py."""
import os, sys, math
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import content as C
import wbbuilder as W
from openpyxl import Workbook
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter

OUT = sys.argv[1] if len(sys.argv) > 1 else "out"

# ---------------------------------------------------------------- helpers ----
def steps_sheet(wb, title, subtitle, hint):
    """Guided click-path sheet. Same look as the follow-along, custom hint."""
    ws = wb.active if wb.active.title == "Sheet" else wb.create_sheet()
    ws.title = "steps"; ws.sheet_properties.tabColor = W.RED[2:]
    ws.sheet_view.showGridLines = False
    for col, w in {"A": 2.5, "B": 15, "C": 74, "D": 2}.items():
        ws.column_dimensions[col].width = w
    ws["B1"] = title; ws["B1"].font = W.F_TITLE; ws.row_dimensions[1].height = 28
    ws.merge_cells("B2:C2"); ws["B2"] = subtitle; ws["B2"].font = W.F_SUB; ws["B2"].alignment = W.WRAP
    ws.merge_cells("B3:C3"); ws["B3"] = hint; ws["B3"].font = W.F_HINTLINE
    ws["B3"].alignment = W.WRAP; ws.row_dimensions[3].height = 28
    return ws, 5

def block(ws, row, label, desc, lines, take=None):
    row = W.step(ws, row, label, desc)
    row = W.code_row(ws, row, "steps", "\n".join(lines))
    row = W.spacer(ws, row)
    if take: row = W.takeaway(ws, row, take)
    return row + 1

def sheet_plain(wb, name, columns, rows, widths=None):
    ws = wb.create_sheet(name)
    ws.append(list(columns))
    for c in ws[1]: c.font = W.F_LABEL
    for r in rows: ws.append(list(r))
    for i, col in enumerate(columns, 1):
        ws.column_dimensions[get_column_letter(i)].width = (widths or {}).get(col, max(12, min(26, len(str(col)) + 6)))
    ws.freeze_panes = "A2"
    return ws

def sheet_table(wb, name, columns, rows, tablename):
    return W.data_sheet(wb, name, columns, rows, tablename)

def save(wb, day, kind):
    d = os.path.join(OUT, f"Day {day}"); os.makedirs(d, exist_ok=True)
    p = os.path.join(d, f"Guerrilla AI - Day {day} - {kind}.xlsx"); wb.save(p); print("wrote", p)

HINT_CLICK = ("Follow the numbered steps below live with George. Everything happens in this "
              "workbook. Nothing to install, nothing to license.")

# --------------------------------------------------------------- the data ----
main_raw = C.junked(C.make_rows())
june = C.june_rows()
raw_plus_name = [r + [C.rep_name(r[2]) if r[2] else ""] for r in main_raw]
no_junk = [r for r in raw_plus_name if r[0] and not str(r[0]).startswith("TOTAL")]
clean_all = C.clean_rows(main_raw) + C.clean_rows(june)
summary = C.region_summary(clean_all)

# ------------------------------------------------------------------ Day 1 ----
def day1(kind):
    wb = Workbook()
    done = kind == "Solution"
    ws, row = steps_sheet(wb, "Day 1: The AI already hiding in your Excel",
                          "Flash Fill, Analyze Data, and Recommended PivotTables on the "
                          f"{C.COMPANY} export.", HINT_CLICK)
    row = block(ws, row, "Demo 1 of 2: Flash Fill",
                "One typed example teaches Excel the pattern.", C.D1_FLASHFILL)
    row = block(ws, row, "Demo 2 of 2: Analyze Data",
                "A natural-language question against your own data.", C.D1_ANALYZE,
                take=C.DAYS[1]["takeaway"])
    cols = C.RAW_COLUMNS + (["Rep Name"] if done else [])
    rows = raw_plus_name if done else main_raw
    sheet_plain(wb, "raw-export", cols, rows)
    save(wb, 1, kind)

# ------------------------------------------------------------------ Day 2 ----
def day2(kind):
    wb = Workbook()
    done = kind == "Solution"
    ws, row = steps_sheet(wb, "Day 2: Make your data AI-readable",
                          "Ctrl+T, a real table name, and structured references.", HINT_CLICK)
    row = block(ws, row, "Demo 1 of 2: Convert to a table",
                "Junk rows out, Ctrl+T, and a name that means something.", C.D2_TABLE)
    row = block(ws, row, "Demo 2 of 2: Rerun Monday's question",
                "Analyze Data on a table versus a raw range.", C.D2_RERUN,
                take=C.DAYS[2]["takeaway"])
    if done:
        row = W.step(ws, row, "Quick answers", "Structured references written live:")
        ws[f"C{row}"] = "=SUM(SalesData[Amount])"; ws[f"C{row}"].font = W.F_CODE
        ws[f"B{row}"] = "total sales"; ws[f"B{row}"].font = W.F_LABEL; row += 1
        ws[f"C{row}"] = "=COUNTA(SalesData[Order ID])"; ws[f"C{row}"].font = W.F_CODE
        ws[f"B{row}"] = "order count"; ws[f"B{row}"].font = W.F_LABEL
        sheet_table(wb, "raw-export", C.RAW_COLUMNS + ["Rep Name"], no_junk, "SalesData")
    else:
        sheet_plain(wb, "raw-export", C.RAW_COLUMNS + ["Rep Name"], raw_plus_name)
    save(wb, 2, kind)

# ------------------------------------------------------------------ Day 3 ----
def day3(kind):
    wb = Workbook()
    done = kind == "Solution"
    ws, row = steps_sheet(wb, "Day 3: Clean data so AI doesn't lie to you",
                          "Power Query records the cleanup once. Refresh replays it forever.",
                          HINT_CLICK)
    row = block(ws, row, "Demo 1 of 2: Record the cleanup",
                "Six recorded steps, no formulas, no code.", C.D3_TRANSFORMS)
    row = block(ws, row, "Demo 2 of 2: The Refresh moment",
                "New month, one click.", C.D3_REFRESH, take=C.DAYS[3]["takeaway"])
    if done:
        row = W.note_row(ws, row, "This solution workbook shows the loaded result, with June "
                         "pasted into RawExport and SalesClean refreshed. The query itself is "
                         "recorded live in the demo; the steps above are the full recipe.")
        sheet_table(wb, "raw-export", C.RAW_COLUMNS, main_raw + june, "RawExport")
        sheet_table(wb, "SalesClean", C.CLEAN_COLUMNS, clean_all, "SalesClean")
    else:
        sheet_table(wb, "raw-export", C.RAW_COLUMNS, main_raw, "RawExport")
        sheet_plain(wb, "june-rows", C.RAW_COLUMNS, june)
    save(wb, 3, kind)

# ------------------------------------------------------------------ Day 4 ----
def day4(kind):
    wb = Workbook()
    done = kind == "Solution"
    if done:
        ws = wb.active; ws.title = "solution-code"; ws.sheet_properties.tabColor = W.RED[2:]
        ws.sheet_view.showGridLines = False
        for col, w in {"A": 2.5, "B": 15, "C": 74, "D": 2}.items():
            ws.column_dimensions[col].width = w
        ws["B1"] = "Day 4 solution code, in run order"; ws["B1"].font = W.F_TITLE
        ws.row_dimensions[1].height = 28
        row = 3
        for i, (label, lines) in enumerate([("Block 1", C.D4_BLOCK1), ("Block 2", C.D4_BLOCK2),
                                            ("Block 3", C.D4_BLOCK3)], 1):
            row = W.step(ws, row, f"{label}: dry-run this into a PY() cell")
            row = W.code_row(ws, row, "PY()", "\n".join(lines))
            row = W.berth(ws, row, 4)
    else:
        ws = W.follow_sheet(wb, "Day 4: A data science lab behind =PY()",
                            "pandas and seaborn are already inside Excel. "
                            "Free chatbots write the code. You supply the judgment.")
        row = 6
        row = W.typealong(ws, row)
        row = W.step(ws, row, "Step 1: Load and describe",
                     "xl() hands the SalesClean table to pandas as a DataFrame.")
        row = W.example(ws, row, "PY()", C.D4_BLOCK1, clear=12)
        row = W.step(ws, row, "Step 2: Sales by region",
                     "One groupby answers what the pivot took clicks to build.")
        row = W.example(ws, row, "PY()", C.D4_BLOCK2, clear=10)
        row = W.step(ws, row, "Step 3: One chart",
                     "Ask a free chatbot for this code, paste it here, run it.")
        row = W.example(ws, row, "PY()", C.D4_BLOCK3, clear=16)
        row = W.takeaway(ws, row, C.DAYS[4]["takeaway"])
    sheet_table(wb, "SalesClean", C.CLEAN_COLUMNS, clean_all, "SalesClean")
    save(wb, 4, kind)

# ------------------------------------------------------------------ Day 5 ----
def day5(kind):
    wb = Workbook()
    done = kind == "Solution"
    structured = C.d5_structured(summary)
    if done:
        ws, row = steps_sheet(wb, "Day 5 solution: the structured prompt",
                              "The exact Markdown sent to the chatbot, and the table "
                              "that came back.", "Compare this against the blob prompt "
                              "from the live session.")
        row = W.step(ws, row, "The structured prompt (Markdown)")
        row = W.code_row(ws, row, "prompt.md", "\n".join(structured))
        pasted = wb.create_sheet("pasted-back")
        pasted.append(["Region", "Total Sales", "Orders", "Share of Sales"])
        for c in pasted[1]: c.font = W.F_LABEL
        grand = sum(t for _, t, _ in summary)
        for reg, tot, cnt in summary:
            pasted.append([reg, tot, cnt, round(tot / grand, 3)])
        for i in range(1, 5): pasted.column_dimensions[get_column_letter(i)].width = 16
    else:
        ws, row = steps_sheet(wb, "Day 5: Markdown, the language AI speaks",
                              "A structured prompt, then the Excel round-trip.", HINT_CLICK)
        row = block(ws, row, "Demo 1 of 2: Structure beats the blob",
                    "Same request, two prompts, very different answers.",
                    C.D5_BLOB + [""] + structured)
        row = block(ws, row, "Demo 2 of 2: The Excel round-trip",
                    "Markdown tables move data without connectors or licenses.",
                    C.D5_ROUNDTRIP, take=C.DAYS[5]["takeaway"])
    sheet_table(wb, "region-summary", ["Region", "Total Sales", "Orders"],
                [[r, t, c] for r, t, c in summary], "RegionSummary")
    save(wb, 5, kind)

for f in (day1, day2, day3, day4, day5):
    f("Demo Start"); f("Solution")
print("workbooks done")
